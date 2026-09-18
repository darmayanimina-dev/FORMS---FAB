from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import JSONResponse, Response
from pypdf import PdfReader
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import re, io, json
from datetime import datetime

app = FastAPI()

def parse_opex(raw_text):
    rows = []
    lines = raw_text.splitlines()

    # Regex pendukung
    date_pattern = re.compile(r'(\d{2}/\d{2}/\d{4})')
    account_code_pattern = re.compile(r'\b(7\d{9})\b')

    # Tangkap periode
    period_match = re.search(r'Periode\s*(?:YTD)?\s*[:|]?\s*([A-Za-z]{3}-\d{2,4})', raw_text, re.IGNORECASE)
    detected_period = period_match.group(1) if period_match else datetime.now().strftime("%b-%y")

    # Daftar kata sampah bawaan header kolom/instansi
    trash_words = [
        "Reg. Retail Sales Hea", "Reg.Retail Sales Hea", "Reg. Retail Sales Heal",
        "Reg.Retail Sales Heal", "PT ASF", "PT", "ASF", "0000", "10000", "None", "|"
    ]

    for line in lines:
        line_str = line.strip()
        if not line_str or "DETAIL REALISASI OPEX" in line_str or "TOTAL:" in line_str:
            continue

        acc_match = account_code_pattern.search(line_str)
        if not acc_match:
            continue

        code_account = acc_match.group(1)
        
        # Tangkap tanggal
        date_match = date_pattern.search(line_str)
        trans_date = date_match.group(1) if date_match else "-"

        # Potong string di kode akun
        parts = line_str.split(code_account)
        before_code = parts[0]
        after_code = parts[1] if len(parts) > 1 else ""

        # Cari saldo di after_code (mendukung angka minus & koma)
        amt_matches = list(re.finditer(r'(-?[\d]{1,3}(?:,\d{3})+(?:\.\d+)?|-?\d{4,})', after_code))
        saldo = 0.0
        description = ""
        account_name = "UNKNOWN"

        if amt_matches:
            target_amt = amt_matches[0]
            try:
                saldo = float(target_amt.group(1).replace(',', ''))
            except ValueError:
                saldo = 0.0

            description = after_code[target_amt.end():].strip().lstrip('|').strip()
            
            # Ambil account name di antara kode akun dan nominal
            acc_cand = after_code[:target_amt.start()]
            for t in trash_words:
                acc_cand = acc_cand.replace(t, "")
            account_name = acc_cand.strip()
        else:
            # Jika nominal tidak ditemukan di after_code, cek before_code
            amt_matches_before = list(re.finditer(r'(-?[\d]{1,3}(?:,\d{3})+(?:\.\d+)?|-?\d{4,})', before_code))
            if amt_matches_before:
                target_amt = amt_matches_before[-1]
                try:
                    saldo = float(target_amt.group(1).replace(',', ''))
                except ValueError:
                    saldo = 0.0
            description = after_code.strip().lstrip('|').strip()

        # Bersihkan nama akun jika masih kosong
        if not account_name or account_name == "UNKNOWN":
            cand = before_code
            for t in trash_words:
                cand = cand.replace(t, "")
            if date_match:
                cand = cand.replace(trans_date, "")
            account_name = cand.replace("|", "").strip() or "BIAYA LAIN"

        # Tangkap batch name secara ringkas
        batch_cand = before_code.replace("PT ASF", "").replace("PT", "").replace("ASF", "")
        if date_match:
            batch_cand = batch_cand.replace(trans_date, "")
        batch_name = batch_cand.replace("|", "").strip() or "-"

        rows.append({
            "Batch Name": batch_name,
            "Trans. Date": trans_date,
            "Code Account": code_account,
            "Account Name": account_name,
            "Saldo": saldo,
            "Description": description
        })

    df = pd.DataFrame(rows)
    return df, detected_period

@app.post("/api/process")
async def process_file(file: UploadFile = File(...)):
    try:
        contents = await file.read()
        raw_text = ""
        if file.filename.lower().endswith(".pdf"):
            reader = PdfReader(io.BytesIO(contents))
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    raw_text += "\n" + t
        else:
            raw_text = contents.decode("utf-8", errors="ignore")

        df, period = parse_opex(raw_text)
        
        if df.empty:
            return JSONResponse(
                status_code=400, 
                content={"error": "Format data laporan tidak terdeteksi. Pastikan file berisi kode akun OPEX (10 digit)."}
            )

        # Kelompokkan per Akun
        summary = (
            df.groupby(["Code Account", "Account Name"], as_index=False)["Saldo"]
            .sum()
            .rename(columns={"Saldo": "Total Saldo"})
            .sort_values(by="Total Saldo", ascending=False)
        )

        grand_total = float(summary["Total Saldo"].sum())

        return {
            "period": period,
            "grand_total": grand_total,
            "account_count": len(summary),
            "summary": summary.to_dict(orient="records"),
            "detail": df.to_dict(orient="records")
        }
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": f"Terjadi kesalahan teknis: {str(e)}"})

@app.post("/api/download/excel")
async def download_excel(data_json: str = Form(...), period: str = Form(...)):
    rows = json.loads(data_json)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detail OPEX"

    headers = ["Batch Name", "Trans. Date", "Code Account", "Account Name", "Saldo", "Description"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx in range(1, 7):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, r in enumerate(rows):
        ws.append([r["Batch Name"], r["Trans. Date"], r["Code Account"], r["Account Name"], r["Saldo"], r["Description"]])
        curr = r_idx + 2
        for c_idx in range(1, 7):
            c = ws.cell(row=curr, column=c_idx)
            c.border = thin_border
            if c_idx == 5:
                c.number_format = '#,##0'
                c.alignment = Alignment(horizontal="right")
            elif c_idx in [2, 3]:
                c.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Detail Realisasi OPEX - {period}.xlsx"'}
    )

@app.post("/api/download/pdf")
async def download_pdf(summary_json: str = Form(...), grand_total: float = Form(...), period: str = Form(...)):
    rows = json.loads(summary_json)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=13, textColor=colors.HexColor('#1F4E78'), spaceAfter=3, alignment=1)
    subtitle_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontSize=9, textColor=colors.gray, spaceAfter=14, alignment=1)

    story.append(Paragraph(f"<b>REKAP REALISASI OPEX - {period.upper()}</b>", title_style))
    story.append(Paragraph("FORMS | Fast, Accurate, and Accessible OPEX Insights", subtitle_style))

    table_data = [["Code Account", "Account Name", "Total Saldo"]]
    cell_style = ParagraphStyle('CellNorm', parent=styles['Normal'], fontSize=8, leading=10)

    rows = sorted(rows, key=lambda x: str(x["Code Account"]))
    for r in rows:
        table_data.append([
            str(r["Code Account"]),
            Paragraph(str(r["Account Name"]), cell_style),
            f"Rp {r['Total Saldo']:,.0f}"
        ])

    table_data.append(["TOTAL SEMUA", "", f"Rp {grand_total:,.0f}"])

    t = Table(table_data, colWidths=[100, 280, 150])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 1), (2, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor('#CCCCCC')),
        ('SPAN', (0, -1), (1, -1)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#F2F4F7')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, -1), (-1, -1), 9),
        ('LINEABOVE', (0, -1), (-1, -1), 1.5, colors.HexColor('#1F4E78')),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))

    story.append(t)
    doc.build(story)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="Rekap Realisasi OPEX - {period}.pdf"'}
    )
